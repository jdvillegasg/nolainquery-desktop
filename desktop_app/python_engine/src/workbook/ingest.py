"""Open workbooks and materialize importable tables."""

from __future__ import annotations

import io
import os
import shutil
from pathlib import Path
from typing import Iterable, List, Optional, Set

import pandas as pd

from .cache import (
    catalog_is_valid,
    catalog_cache_dir,
    load_cached_catalog,
    save_catalog,
)
from .config import (
    DEFAULT_INGEST_CONFIG,
    SourceKind,
    SUPPORTED_SOURCE_EXTENSIONS,
    TableRole,
    WorkbookIngestConfig,
)
from .identifiers import materialized_filename, unique_table_id
from .models import WorkbookCatalog, WorkbookTable
from .table_detection import detect_table_in_excel_sheet


class WorkbookIngestError(ValueError):
    """Raised when a workbook cannot be opened or has no importable tables."""


def _resolve_source_kind(source_path: str) -> SourceKind:
    suffix = Path(source_path).suffix.lower()
    for kind, extensions in SUPPORTED_SOURCE_EXTENSIONS.items():
        if suffix in extensions:
            return kind
    raise WorkbookIngestError(f"Unsupported file type: {suffix or '(none)'}")


def _excel_hidden_sheet_names(source_path: str) -> Set[str]:
    suffix = Path(source_path).suffix.lower()
    if suffix != ".xlsx":
        return set()
    try:
        from openpyxl import load_workbook
    except ImportError:
        return set()
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        hidden = set()
        for sheet in workbook.worksheets:
            state = getattr(sheet, "sheet_state", "visible")
            if state in {"hidden", "veryHidden"}:
                hidden.add(sheet.title)
        return hidden
    finally:
        workbook.close()


def _stringify_series(series: pd.Series) -> pd.Series:
    return series.map(lambda value: None if pd.isna(value) else str(value))


def _prepare_frame_for_parquet(frame: pd.DataFrame) -> pd.DataFrame:
    """Stringify object columns so mixed Excel types can be written as parquet."""
    converted: dict[str, pd.Series] = {}
    for column in frame.columns:
        if pd.api.types.is_object_dtype(frame[column]):
            converted[column] = _stringify_series(frame[column])
    if not converted:
        return frame
    return frame.assign(**converted)


def _write_parquet_safe(frame: pd.DataFrame, destination: Path) -> None:
    """Materialize a frame to parquet, stringifying columns that break inference."""
    safe = _prepare_frame_for_parquet(frame)
    try:
        safe.to_parquet(destination, index=False)
        return
    except Exception:
        pass
    for column in safe.columns:
        probe = io.BytesIO()
        try:
            safe[[column]].to_parquet(probe, index=False)
        except Exception:
            safe[column] = _stringify_series(safe[column])
    safe.to_parquet(destination, index=False)


def _materialize_dataframe(
    frame: pd.DataFrame,
    destination: Path,
    config: WorkbookIngestConfig,
) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if config.materialized_format == "parquet":
        try:
            _prepare_frame_for_parquet(frame).to_parquet(destination, index=False)
        except Exception:
            _write_parquet_safe(frame, destination)
        return
    if config.materialized_format == "csv":
        frame.to_csv(destination, index=False)
        return
    raise WorkbookIngestError(
        f"Unsupported materialized format: {config.materialized_format}"
    )


def _copy_or_materialize_single_file(
    source_path: str,
    cache_dir: Path,
    table_id: str,
    config: WorkbookIngestConfig,
) -> tuple[pd.DataFrame, Path]:
    suffix = Path(source_path).suffix.lower()
    if suffix == ".csv":
        destination = cache_dir / materialized_filename(table_id, config.materialized_format)
        frame = pd.read_csv(source_path)
        _materialize_dataframe(frame, destination, config)
        return frame, destination
    if suffix == ".parquet":
        destination = cache_dir / materialized_filename(table_id, ".parquet")
        shutil.copy2(source_path, destination)
        return pd.read_parquet(source_path), destination
    raise WorkbookIngestError(f"Cannot materialize single-file source: {suffix}")


def _select_default_active_table(
    tables: List[WorkbookTable],
    config: WorkbookIngestConfig,
) -> str:
    importable = [table for table in tables if table.role == TableRole.TABLE]
    if not importable:
        raise WorkbookIngestError("No tabular sheets found in this workbook.")
    if config.default_active_table == "first_visible":
        return importable[0].id
    importable.sort(
        key=lambda table: (
            table.row_count or 0,
            table.column_count or 0,
            table.sheet_name.lower(),
        ),
        reverse=True,
    )
    return importable[0].id


def _ingest_excel_workbook(
    source_path: str,
    cache_dir: Path,
    config: WorkbookIngestConfig,
    force_include_sheets: Optional[Iterable[str]] = None,
) -> WorkbookCatalog:
    force_set = {name.strip() for name in (force_include_sheets or []) if name.strip()}
    hidden_sheets = _excel_hidden_sheet_names(source_path)
    excel_file = pd.ExcelFile(source_path)
    used_ids: set[str] = set()
    tables: List[WorkbookTable] = []

    for sheet_name in excel_file.sheet_names:
        table_id = unique_table_id(sheet_name, used_ids)
        force_include = sheet_name in force_set

        if (
            not config.include_hidden_sheets
            and sheet_name in hidden_sheets
            and not force_include
        ):
            tables.append(
                WorkbookTable(
                    id=table_id,
                    sheet_name=sheet_name,
                    role=TableRole.SKIPPED,
                    skip_reason="hidden_sheet",
                )
            )
            continue

        try:
            detection = detect_table_in_excel_sheet(
                excel_file,
                sheet_name,
                config,
                source_path=source_path,
                force_include=force_include,
            )
        except Exception:
            tables.append(
                WorkbookTable(
                    id=table_id,
                    sheet_name=sheet_name,
                    role=TableRole.SKIPPED,
                    skip_reason="read_error",
                )
            )
            continue

        if not detection.is_table or detection.dataframe is None:
            tables.append(
                WorkbookTable(
                    id=table_id,
                    sheet_name=sheet_name,
                    role=TableRole.SKIPPED,
                    skip_reason=(
                        detection.skip_reason.value
                        if detection.skip_reason
                        else "unstructured"
                    ),
                    header_row=detection.header_row,
                )
            )
            continue

        destination = cache_dir / materialized_filename(
            table_id, config.materialized_format
        )
        frame = detection.dataframe
        _materialize_dataframe(frame, destination, config)
        tables.append(
            WorkbookTable(
                id=table_id,
                sheet_name=sheet_name,
                role=TableRole.TABLE,
                materialized_path=str(destination),
                row_count=int(len(frame)),
                column_count=int(frame.shape[1]),
                header_row=detection.header_row,
            )
        )

    active_table_id = _select_default_active_table(tables, config)
    return WorkbookCatalog(
        source_path=os.path.abspath(source_path),
        source_kind=SourceKind.EXCEL,
        source_name=Path(source_path).name,
        active_table_id=active_table_id,
        tables=tables,
        cache_dir=str(cache_dir),
        ingest_version=config.ingest_version,
    )


def _ingest_single_table_source(
    source_path: str,
    source_kind: SourceKind,
    cache_dir: Path,
    config: WorkbookIngestConfig,
) -> WorkbookCatalog:
    table_id = unique_table_id(Path(source_path).stem, set())
    frame, materialized = _copy_or_materialize_single_file(
        source_path, cache_dir, table_id, config
    )
    tables = [
        WorkbookTable(
            id=table_id,
            sheet_name=Path(source_path).stem,
            role=TableRole.TABLE,
            materialized_path=str(materialized),
            row_count=int(len(frame)),
            column_count=int(frame.shape[1]),
        )
    ]
    return WorkbookCatalog(
        source_path=os.path.abspath(source_path),
        source_kind=source_kind,
        source_name=Path(source_path).name,
        active_table_id=table_id,
        tables=tables,
        cache_dir=str(cache_dir),
        ingest_version=config.ingest_version,
    )


def _resolve_active_table_id(
    catalog: WorkbookCatalog,
    *,
    active_table_id: Optional[str] = None,
    active_sheet_name: Optional[str] = None,
) -> str:
    if active_table_id:
        if not any(table.id == active_table_id for table in catalog.tables):
            raise WorkbookIngestError(
                f"Unknown table id for workbook: {active_table_id}"
            )
        return active_table_id
    if active_sheet_name:
        for table in catalog.importable_tables:
            if table.sheet_name == active_sheet_name:
                return table.id
        raise WorkbookIngestError(
            f"No importable table found for sheet: {active_sheet_name}"
        )
    return catalog.active_table_id


def open_workbook(
    source_path: str,
    *,
    active_table_id: Optional[str] = None,
    active_sheet_name: Optional[str] = None,
    force_include_sheets: Optional[Iterable[str]] = None,
    config: WorkbookIngestConfig = DEFAULT_INGEST_CONFIG,
    use_cache: bool = True,
) -> WorkbookCatalog:
    normalized = os.path.abspath(source_path)
    if not os.path.isfile(normalized):
        raise WorkbookIngestError(f"File not found: {source_path}")

    if use_cache and not force_include_sheets:
        cached = load_cached_catalog(normalized)
        if cached and catalog_is_valid(cached):
            cached.active_table_id = _resolve_active_table_id(
                cached,
                active_table_id=active_table_id,
                active_sheet_name=active_sheet_name,
            )
            return cached

    cache_dir = catalog_cache_dir(normalized)
    source_kind = _resolve_source_kind(normalized)
    if source_kind == SourceKind.EXCEL:
        catalog = _ingest_excel_workbook(
            normalized,
            cache_dir,
            config,
            force_include_sheets=force_include_sheets,
        )
    else:
        catalog = _ingest_single_table_source(
            normalized, source_kind, cache_dir, config
        )

    catalog.active_table_id = _resolve_active_table_id(
        catalog,
        active_table_id=active_table_id,
        active_sheet_name=active_sheet_name,
    )

    save_catalog(catalog)
    return catalog
