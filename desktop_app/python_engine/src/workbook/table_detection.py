"""Heuristics to classify Excel sheets as tabular vs non-table."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, List, Optional, Sequence, Tuple

import pandas as pd

from .config import SkipReason, WorkbookIngestConfig


@dataclass(frozen=True)
class TableDetectionResult:
    is_table: bool
    skip_reason: Optional[SkipReason] = None
    header_row: Optional[int] = None
    dataframe: Optional[pd.DataFrame] = None
    fill_density: Optional[float] = None
    header_confidence: Optional[float] = None


def _is_empty(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _trim_bounding_box(raw: pd.DataFrame) -> pd.DataFrame:
    if raw.empty:
        return raw
    mask = raw.apply(lambda series: series.map(lambda value: not _is_empty(value)))
    if not mask.values.any():
        return raw.iloc[0:0, 0:0]
    row_keep = mask.any(axis=1)
    col_keep = mask.any(axis=0)
    return raw.loc[row_keep, col_keep].reset_index(drop=True)


def _fill_density(raw: pd.DataFrame) -> float:
    if raw.empty:
        return 0.0
    total = raw.shape[0] * raw.shape[1]
    if total == 0:
        return 0.0
    filled = sum(not _is_empty(value) for value in raw.to_numpy().ravel())
    return filled / total


_DATE_VALUE = re.compile(
    r"^\d{4}-\d{2}-\d{2}(?:[ T]\d{2}:\d{2}:\d{2})?"
)
_PERIOD_VALUE = re.compile(
    r"^(?:Q[1-4]|FY)\s*'?\d{2,4}$",
    re.IGNORECASE,
)


def _looks_like_data_value(value: Any) -> bool:
    """True for numbers, dates, and period labels — not field names."""
    if _is_empty(value) or isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    if hasattr(value, "year") and hasattr(value, "month"):
        return True
    text = str(value).strip()
    return bool(_DATE_VALUE.match(text) or _PERIOD_VALUE.match(text))


def _looks_like_header_cell(value: Any) -> bool:
    if _is_empty(value) or _looks_like_data_value(value):
        return False
    return True


def _is_key_value_row(row: Any) -> bool:
    """A label plus a value is metadata, not a table header."""
    cells = [value for value in row if not _is_empty(value)]
    if len(cells) != 2:
        return False
    return _looks_like_header_cell(cells[0]) and _looks_like_data_value(cells[1])


def _header_confidence(
    raw: pd.DataFrame,
    header_row: int,
    min_header_cells: int = 2,
    consistency_rows: int = 5,
) -> float:
    if raw.empty or header_row >= len(raw):
        return 0.0
    header = raw.iloc[header_row]
    if _is_key_value_row(header):
        return 0.0
    header_cells = [cell for cell in header if _looks_like_header_cell(cell)]
    if len(header_cells) < min_header_cells:
        return 0.0
    non_empty_header = len(header_cells)
    width = max(int(raw.shape[1]), 1)
    header_score = min(1.0, non_empty_header / width)
    data_rows = raw.iloc[header_row + 1 : header_row + 1 + max(consistency_rows, 1)]
    if data_rows.empty:
        return header_score * 0.5
    consistency = 0.0
    for column_index in range(raw.shape[1]):
        column_values = [
            value
            for value in data_rows.iloc[:, column_index].tolist()
            if not _is_empty(value)
        ]
        if not column_values:
            continue
        consistency += 1.0
    consistency_score = consistency / width
    return (header_score + consistency_score) / 2


def _find_header_row(raw: pd.DataFrame, config: WorkbookIngestConfig) -> Optional[int]:
    scan_limit = min(config.header_scan_rows, len(raw))
    best_row: Optional[int] = None
    best_score = -1.0
    for row_index in range(scan_limit):
        score = _header_confidence(
            raw,
            row_index,
            config.min_columns,
            config.header_consistency_rows,
        )
        if score > best_score:
            best_score = score
            best_row = row_index
    if best_row is None or best_score < config.min_header_confidence:
        return None
    return best_row


def _unnamed_ratio(columns: Sequence[Any]) -> float:
    if len(columns) == 0:
        return 1.0
    unnamed = 0
    for column in columns:
        name = str(column)
        if name.startswith("Unnamed:") or not name.strip():
            unnamed += 1
    return unnamed / len(columns)


def _normalize_columns(columns: Sequence[Any]) -> List[str]:
    normalized: List[str] = []
    seen: dict[str, int] = {}
    for index, column in enumerate(columns):
        label = str(column).strip() if not _is_empty(column) else f"column_{index + 1}"
        if not label or label.startswith("Unnamed:"):
            label = f"column_{index + 1}"
        count = seen.get(label, 0)
        if count:
            label = f"{label}_{count + 1}"
        seen[label] = count + 1
        normalized.append(label)
    return normalized


def _extract_table_frame(
    raw: pd.DataFrame, header_row: int
) -> pd.DataFrame:
    header = raw.iloc[header_row].tolist()
    data = raw.iloc[header_row + 1 :].copy()
    data.columns = _normalize_columns(header)
    data = data.dropna(how="all").reset_index(drop=True)
    data = data.loc[:, ~data.apply(lambda series: series.map(_is_empty).all())]
    return data


def _finalize_headered_frame(frame: pd.DataFrame) -> pd.DataFrame:
    data = frame.copy()
    data.columns = _normalize_columns(data.columns)
    data = data.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
    return data


def _pad_sample_to_used_extent(
    sample: pd.DataFrame, row_count: int, col_count: int
) -> pd.DataFrame:
    """Expand a pandas sample to the sheet used range so sparsity is not inflated."""
    if row_count <= 0 or col_count <= 0:
        return sample
    if len(sample) >= row_count and sample.shape[1] >= col_count:
        return sample.iloc[:row_count, :col_count]
    return sample.reindex(index=range(row_count), columns=range(col_count))


def _excel_sheet_used_extent(
    source_path: str, sheet_name: str, sample_rows: int
) -> Optional[Tuple[int, int]]:
    suffix = Path(source_path).suffix.lower()
    if suffix != ".xlsx":
        return None
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        used_rows = min(int(worksheet.max_row or 0), sample_rows)
        used_cols = int(worksheet.max_column or 0)
        if used_rows <= 0 or used_cols <= 0:
            return None
        return used_rows, used_cols
    finally:
        workbook.close()


def detect_table_in_excel_sheet(
    excel_file: pd.ExcelFile,
    sheet_name: str,
    config: WorkbookIngestConfig,
    *,
    source_path: Optional[str] = None,
    force_include: bool = False,
) -> TableDetectionResult:
    """Detect a table from a sample, then load the full sheet only when needed."""
    sample_rows = config.header_sample_rows()
    sample = pd.read_excel(
        excel_file,
        sheet_name=sheet_name,
        header=None,
        nrows=sample_rows,
    )
    used_extent = (
        _excel_sheet_used_extent(source_path, sheet_name, sample_rows)
        if source_path
        else None
    )
    if used_extent is not None:
        sample = _pad_sample_to_used_extent(sample, *used_extent)
    detection = detect_table(sample, config, force_include=force_include)
    if not detection.is_table:
        return detection

    header_row = detection.header_row if detection.header_row is not None else 0
    if len(sample) < sample_rows:
        return detection

    frame = pd.read_excel(excel_file, sheet_name=sheet_name, header=header_row)
    frame = _finalize_headered_frame(frame)
    return TableDetectionResult(
        is_table=True,
        header_row=header_row,
        dataframe=frame,
        fill_density=detection.fill_density,
        header_confidence=detection.header_confidence,
    )


def detect_table(
    raw: pd.DataFrame,
    config: WorkbookIngestConfig,
    *,
    force_include: bool = False,
) -> TableDetectionResult:
    trimmed = _trim_bounding_box(raw)
    if trimmed.empty:
        return TableDetectionResult(
            is_table=False,
            skip_reason=SkipReason.EMPTY_SHEET,
            fill_density=0.0,
        )

    # Measure sparsity on the full sample grid. Trimming first collapses empty
    # padding and inflates density for cover pages and KPI callout layouts.
    density = _fill_density(raw)
    if not force_include and density < config.min_fill_density:
        return TableDetectionResult(
            is_table=False,
            skip_reason=SkipReason.SPARSE_NON_TABLE,
            fill_density=density,
        )

    # Header detection also runs on the untrimmed grid so deep headers are not
    # pulled up by removed blank rows above the table.
    header_row = _find_header_row(raw, config)
    if header_row is None and not force_include:
        return TableDetectionResult(
            is_table=False,
            skip_reason=SkipReason.LOW_HEADER_CONFIDENCE,
            fill_density=density,
        )
    if header_row is None:
        header_row = 0

    confidence = _header_confidence(raw, header_row, config.min_columns)
    frame = _extract_table_frame(raw, header_row)

    if frame.shape[1] < config.min_columns and not force_include:
        return TableDetectionResult(
            is_table=False,
            skip_reason=SkipReason.TOO_FEW_COLUMNS,
            fill_density=density,
            header_confidence=confidence,
        )
    if len(frame) < config.min_data_rows and not force_include:
        return TableDetectionResult(
            is_table=False,
            skip_reason=SkipReason.TOO_FEW_ROWS,
            fill_density=density,
            header_confidence=confidence,
        )

    unnamed_ratio = _unnamed_ratio(frame.columns)
    if unnamed_ratio > config.max_unnamed_column_ratio and not force_include:
        return TableDetectionResult(
            is_table=False,
            skip_reason=SkipReason.LOW_HEADER_CONFIDENCE,
            fill_density=density,
            header_confidence=confidence,
        )

    return TableDetectionResult(
        is_table=True,
        header_row=header_row,
        dataframe=frame,
        fill_density=density,
        header_confidence=confidence,
    )
