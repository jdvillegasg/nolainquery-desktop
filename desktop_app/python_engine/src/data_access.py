"""Bounded data access helpers shared by local API endpoints and preprocessors."""
from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any, List

import pandas as pd
import polars as pl


from src.workbook.config import get_ingest_config


def _preview_row_cap() -> int:
    try:
        return int(get_ingest_config().preview_row_cap)
    except Exception:
        return 100_000


PREVIEW_SAMPLE_ROWS = _preview_row_cap()
MAX_EXCEL_ROWS = PREVIEW_SAMPLE_ROWS


@dataclass(frozen=True)
class PreviewPage:
    columns: List[str]
    rows: List[dict[str, Any]]
    total_rows: int
    sampled_rows: int
    is_sampled: bool


@dataclass(frozen=True)
class DistinctValues:
    values: List[Any]
    total_count: int


def _suffix(source_path: str) -> str:
    return Path(source_path).suffix.lower()


def scan_csv_lazy(source_path: str) -> pl.LazyFrame:
    """Lazily scan CSV with schema inferred from the full file.

    Polars defaults to ``infer_schema_length=100``, which mis-types columns that
    look numeric in the header sample but contain string IDs later (e.g. invoice
    numbers prefixed with ``C`` for cancellations).
    """
    return pl.scan_csv(source_path, infer_schema_length=None)


def lazy_frame(source_path: str) -> pl.LazyFrame:
    """Return a lazy CSV/Parquet scan without materializing source rows."""
    suffix = _suffix(source_path)
    if suffix == ".csv":
        return scan_csv_lazy(source_path)
    if suffix == ".parquet":
        return pl.scan_parquet(source_path)
    raise ValueError("Lazy access is supported only for CSV and Parquet files")


def collect_lazy(lf: pl.LazyFrame) -> pl.DataFrame:
    """Collect with the streaming API supported by the installed Polars version."""
    version = tuple(
        int(part) for part in pl.__version__.split(".")[:2]
    )
    if version >= (1, 25):
        return lf.collect(engine="streaming")
    return lf.collect(streaming=True)


def polars_to_pandas(frame: pl.DataFrame) -> pd.DataFrame:
    """Convert a Polars frame to Pandas without requiring PyArrow at runtime.

    Polars' ``to_pandas()`` imports PyArrow. PyInstaller does not collect that
    dependency from a dynamic import, so the packaged sidecar raises and the UI
    surfaces it as "Unsupported file format".
    """
    try:
        return frame.to_pandas()
    except Exception:
        return pd.DataFrame(
            {name: frame.get_column(name).to_list() for name in frame.columns}
        )


def lazy_row_count(source_path: str) -> int:
    return int(collect_lazy(lazy_frame(source_path).select(pl.len())).item())


def _safe_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if hasattr(value, "item"):
        try:
            return _safe_value(value.item())
        except (TypeError, ValueError):
            pass
    return value


def _safe_records(records: List[dict[str, Any]]) -> List[dict[str, Any]]:
    return [
        {str(column): _safe_value(value) for column, value in row.items()}
        for row in records
    ]


def lazy_preview_page(
    source_path: str,
    page: int,
    page_size: int,
    sample_threshold: int = PREVIEW_SAMPLE_ROWS,
) -> PreviewPage:
    """Read one preview page while keeping CSV/Parquet execution lazy."""
    lf = lazy_frame(source_path)
    columns = list(lf.collect_schema().names())
    total_rows = lazy_row_count(source_path)
    is_sampled = total_rows > sample_threshold

    if is_sampled:
        step = max(1, math.ceil(total_rows / sample_threshold))
        sampled_rows = math.ceil(total_rows / step)
        view = (
            lf.with_row_index("__nolain_row_index")
            .filter((pl.col("__nolain_row_index") % step) == 0)
            .drop("__nolain_row_index")
        )
    else:
        sampled_rows = total_rows
        view = lf

    offset = page * page_size
    frame = collect_lazy(view.slice(offset, page_size))
    return PreviewPage(
        columns=columns,
        rows=_safe_records(frame.to_dicts()),
        total_rows=total_rows,
        sampled_rows=sampled_rows,
        is_sampled=is_sampled,
    )


def lazy_distinct_values(
    source_path: str,
    column: str,
    search: str,
    limit: int,
) -> DistinctValues:
    """Return only the requested distinct values plus an exact filtered count."""
    lf = lazy_frame(source_path)
    if column not in lf.collect_schema().names():
        raise KeyError(column)

    values = lf.select(pl.col(column)).filter(pl.col(column).is_not_null())
    if search:
        values = values.filter(
            pl.col(column)
            .cast(pl.String)
            .str.to_lowercase()
            .str.contains(search.lower(), literal=True)
        )

    total_count = int(
        collect_lazy(values.select(pl.col(column).n_unique().alias("count"))).item()
    )
    limited = collect_lazy(
        values.select(pl.col(column).unique(maintain_order=True)).limit(limit)
    ).get_column(column)
    return DistinctValues(
        values=[_safe_value(value) for value in limited.to_list()],
        total_count=total_count,
    )


def _xlsx_preview_page(source_path: str, page: int, page_size: int) -> PreviewPage:
    """Use openpyxl read-only mode so XLSX previews remain memory bounded."""
    from openpyxl import load_workbook

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        columns = [str(value) if value is not None else f"column_{i}" for i, value in enumerate(header)]
        total_rows = max(0, sheet.max_row - 1)
        is_sampled = total_rows > PREVIEW_SAMPLE_ROWS
        step = max(1, math.ceil(total_rows / PREVIEW_SAMPLE_ROWS)) if is_sampled else 1
        sampled_rows = math.ceil(total_rows / step) if total_rows else 0
        wanted_start = page * page_size
        wanted_end = min(wanted_start + page_size, sampled_rows)
        source_rows = {2 + index * step for index in range(wanted_start, wanted_end)}

        rows: List[dict[str, Any]] = []
        if source_rows:
            for row_number, values in enumerate(
                sheet.iter_rows(
                    min_row=min(source_rows),
                    max_row=max(source_rows),
                    values_only=True,
                ),
                start=min(source_rows),
            ):
                if row_number in source_rows:
                    rows.append(
                        {
                            column: _safe_value(values[i] if i < len(values) else None)
                            for i, column in enumerate(columns)
                        }
                    )
        return PreviewPage(columns, rows, total_rows, sampled_rows, is_sampled)
    finally:
        workbook.close()


def excel_preview_page(source_path: str, page: int, page_size: int) -> PreviewPage:
    """Bound Excel previews, with read-only XLSX and a capped legacy-XLS fallback."""
    if _suffix(source_path) == ".xlsx":
        return _xlsx_preview_page(source_path, page, page_size)

    # Legacy XLS readers do not expose a common streaming API. Cap the fallback.
    excel_file = pd.ExcelFile(source_path)
    frame = pd.read_excel(excel_file, nrows=MAX_EXCEL_ROWS)
    try:
        total_rows = max(0, int(excel_file.book.sheet_by_index(0).nrows) - 1)
    except (AttributeError, TypeError, ValueError):
        total_rows = len(frame)
    is_sampled = total_rows > MAX_EXCEL_ROWS
    bounded = frame
    page_frame = bounded.iloc[page * page_size : (page + 1) * page_size]
    return PreviewPage(
        columns=[str(column) for column in frame.columns],
        rows=_safe_records(page_frame.to_dict(orient="records")),
        total_rows=total_rows,
        sampled_rows=len(bounded),
        is_sampled=is_sampled,
    )


def excel_distinct_values(
    source_path: str,
    column: str,
    search: str,
    limit: int,
) -> DistinctValues:
    """Compute distinct Excel values from a capped sample."""
    try:
        frame = pd.read_excel(
            source_path, usecols=[column], nrows=MAX_EXCEL_ROWS
        )
    except ValueError as exc:
        columns = pd.read_excel(source_path, nrows=0).columns
        if column not in columns:
            raise KeyError(column) from exc
        raise
    if column not in frame.columns:
        raise KeyError(column)
    values = frame[column].dropna()
    if search:
        values = values[values.astype(str).str.contains(search, case=False, regex=False)]
    unique = values.drop_duplicates()
    return DistinctValues(
        values=[_safe_value(value) for value in unique.iloc[:limit].tolist()],
        total_count=int(len(unique)),
    )


def bounded_dataframe(
    source_path: str,
    columns: List[str] | None = None,
    max_rows: int = PREVIEW_SAMPLE_ROWS,
) -> pd.DataFrame:
    """Load a bounded, representative frame for non-contractual heavy analysis."""
    suffix = _suffix(source_path)
    if suffix in (".csv", ".parquet"):
        lf = lazy_frame(source_path)
        if columns:
            lf = lf.select(columns)
        total_rows = int(collect_lazy(lf.select(pl.len())).item())
        if total_rows > max_rows:
            step = math.ceil(total_rows / max_rows)
            lf = (
                lf.with_row_index("__nolain_row_index")
                .filter((pl.col("__nolain_row_index") % step) == 0)
                .drop("__nolain_row_index")
            )
        frame = polars_to_pandas(collect_lazy(lf.limit(max_rows)))
        frame.attrs["source_total_rows"] = total_rows
        return frame

    frame = pd.read_excel(source_path, usecols=columns, nrows=max_rows)
    frame.attrs["source_total_rows"] = len(frame)
    return frame
